"""Müşteri Excel import/export.

Export sütunları : Ad Soyad | Telefon | Adres 1 | Adres 2 | … (dinamik)
Import sütunları : Ad Soyad | Telefon | Adres 1 | Adres 2 | … (dinamik)
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from database import customers_db
from utils.helpers import format_phone, normalize_phone


# Sabit sütunlar (adres sütunları dinamik eklenir)
_BASE_COLUMNS = ["Ad Soyad", "Telefon"]


def export_customers(file_path: str | Path) -> int:
    """Veritabanındaki tüm müşterileri Excel'e yazar. Yazılan satır sayısını döner."""
    file_path = Path(file_path)
    customers = customers_db.list_customers()

    # Önce tüm müşterilerin adreslerini çek, maksimum adres sayısını bul
    cust_addresses: list[tuple[dict, list[dict]]] = []
    max_addr = 0
    for cust in customers:
        addrs = customers_db.list_addresses(cust["id"])
        if len(addrs) > max_addr:
            max_addr = len(addrs)
        cust_addresses.append((cust, addrs))

    max_addr = max(max_addr, 1)  # en az 1 adres sütunu
    columns = list(_BASE_COLUMNS)
    for i in range(1, max_addr + 1):
        columns.append(f"Adres {i}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Müşteriler"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="1976D2", end_color="1976D2", fill_type="solid"
    )
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 2
    for cust, addrs in cust_addresses:
        ws.cell(row=row_idx, column=1, value=cust.get("name", ""))
        ws.cell(row=row_idx, column=2, value=format_phone(cust.get("phone", "")))
        for ai, addr in enumerate(addrs):
            ws.cell(row=row_idx, column=3 + ai, value=addr.get("address_text", ""))
        row_idx += 1

    # Sütun genişlikleri
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 16
    for ci in range(3, 3 + max_addr):
        col_letter = ws.cell(row=1, column=ci).column_letter
        ws.column_dimensions[col_letter].width = 55

    wb.save(file_path)
    return row_idx - 2


def import_customers(file_path: str | Path) -> dict[str, int]:
    """Excel'den müşterileri içe aktarır.

    Sütunlar: Ad Soyad | Telefon | Adres 1 | Adres 2 | …

    Aynı telefondaki müşteri varsa eklemez (skip).
    Tüm adres sütunları okunur ve müşteriye eklenir.
    """
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(file_path)

    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    inserted = 0
    skipped = 0
    failed = 0

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"inserted": 0, "skipped": 0, "failed": 0, "total": 0}

    header = [str(c).strip() if c else "" for c in rows[0]]
    name_idx = _find_column(header, ["Ad Soyad", "Ad", "İsim", "İSİM", "Name"])
    phone_idx = _find_column(header, ["Telefon", "TELEFON", "Phone", "Tel"])

    if name_idx is None or phone_idx is None:
        raise ValueError(
            "Excel dosyasında en az 'Ad Soyad' ve 'Telefon' sütunları olmalı."
        )

    # Tüm adres sütunlarını bul (Adres, Adres 1, Adres 2, ADRES 1, …)
    addr_indices = _find_address_columns(header)

    for row in rows[1:]:
        try:
            name = _cell_text(row, name_idx)
            phone_raw = _cell_text(row, phone_idx)

            if not name or not phone_raw:
                failed += 1
                continue

            phone = normalize_phone(phone_raw)
            if not phone:
                failed += 1
                continue

            if customers_db.get_customer_by_phone(phone):
                skipped += 1
                continue

            cust = customers_db.create_customer(name=name, phone=phone)
            cust_id = cust.get("id")
            if cust_id:
                for ai, aidx in enumerate(addr_indices):
                    addr_text = _cell_text(row, aidx)
                    if addr_text:
                        customers_db.create_address(
                            cust_id,
                            address_text=addr_text,
                            title=f"Adres {ai + 1}",
                            is_default=(ai == 0),
                        )
            inserted += 1
        except Exception:  # noqa: BLE001
            failed += 1

    return {
        "inserted": inserted,
        "skipped": skipped,
        "failed": failed,
        "total": len(rows) - 1,
    }


# ---------------- internal ----------------
def _find_column(header: list[str], aliases: list[str]) -> int | None:
    lowered = [h.strip().lower() for h in header]
    for alias in aliases:
        target = alias.strip().lower()
        if target in lowered:
            return lowered.index(target)
    return None


def _find_address_columns(header: list[str]) -> list[int]:
    """Başlıkta adres sütunlarını sırasıyla bulur.

    Eşleşen kalıplar: Adres, Adres 1, Adres 2, ADRES 1, ADRES 2, Address, …
    Sonuç sütun indekslerinin sıralı listesidir.
    """
    import re
    indices: list[int] = []
    for i, h in enumerate(header):
        h_clean = h.strip().lower()
        if re.match(r"^adres\s*\d*$", h_clean) or re.match(r"^address\s*\d*$", h_clean):
            indices.append(i)
    return indices


def _cell_text(row: tuple[Any, ...], idx: int | None) -> str:
    if idx is None or idx >= len(row):
        return ""
    val = row[idx]
    if val is None:
        return ""
    return str(val).strip()

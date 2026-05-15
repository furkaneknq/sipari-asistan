"""MySQL dump'tan müşterileri Excel'e aktar (isim, telefon, adres1, adres2)."""
import csv
import io
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment

SQL_FILE = Path("/home/furkan/Masaüstü/Proje/Musteriler_20260515.sql")
OUTPUT = Path("/home/furkan/Masaüstü/Proje/Musteriler.xlsx")
ENCODING = "iso-8859-9"

IDX_NAME = 3
IDX_ADDR1 = 4
IDX_PHONE = 7
IDX_ADDR2 = 32


def normalize_phone(raw: str) -> str:
    digits = re.sub(r"\D", "", raw)
    if len(digits) == 12 and digits.startswith("90"):
        digits = "0" + digits[2:]
    elif len(digits) == 11 and digits.startswith("0"):
        pass
    elif len(digits) == 10 and digits.startswith("5"):
        digits = "0" + digits
    return digits


def parse_customers() -> list[dict]:
    content = SQL_FILE.read_bytes().decode(ENCODING, errors="replace")
    blocks = re.findall(r"INSERT INTO .musteriler. VALUES\r?\n(.*?);", content, re.DOTALL)

    customers: list[dict] = []
    seen_phones: set[str] = set()

    for blk in blocks:
        tuples_raw = re.split(r"\),\s*\r?\n\s*\(", blk)
        for t in tuples_raw:
            t = t.strip().lstrip("(").rstrip(")")
            reader = csv.reader(io.StringIO(t), quotechar="'", skipinitialspace=True)
            try:
                parts = next(reader)
            except Exception:
                continue
            if len(parts) < 33:
                continue

            name = parts[IDX_NAME].strip()
            addr1 = parts[IDX_ADDR1].strip()
            phone_raw = parts[IDX_PHONE].strip()
            addr2 = parts[IDX_ADDR2].strip() if parts[IDX_ADDR2].strip() != "NULL" else ""

            phone = normalize_phone(phone_raw)
            if not phone or len(phone) < 10:
                continue
            if not name or name.upper() in ("KASA", "TEST", "SİL", "SIL"):
                continue
            if phone in seen_phones:
                continue
            seen_phones.add(phone)

            customers.append({
                "name": name.upper().replace("\\\\", "/"),
                "phone": phone,
                "adres1": addr1.upper().replace("\\\\", "/") if addr1 and addr1 != "NULL" else "",
                "adres2": addr2.upper().replace("\\\\", "/") if addr2 else "",
            })

    return customers


def write_excel(customers: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Müşteriler"

    # Başlıklar
    headers = ["İSİM", "TELEFON", "ADRES 1", "ADRES 2"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Veriler
    for row_idx, cust in enumerate(customers, 2):
        ws.cell(row=row_idx, column=1, value=cust["name"])
        ws.cell(row=row_idx, column=2, value=cust["phone"])
        ws.cell(row=row_idx, column=3, value=cust["adres1"])
        ws.cell(row=row_idx, column=4, value=cust["adres2"])

    # Sütun genişlikleri
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 60

    wb.save(OUTPUT)
    print(f"Excel kaydedildi: {OUTPUT}")


if __name__ == "__main__":
    print("Parse ediliyor...")
    customers = parse_customers()
    print(f"Toplam: {len(customers)} müşteri ({sum(1 for c in customers if c['adres2'])} adet 2 adresli)")
    write_excel(customers)
